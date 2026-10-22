"""
Cruce de las tarifas negociadas con cada cliente contra el costo SICETAC.

SICETAC publica el costo de referencia de UN VIAJE por ruta, configuración del
vehículo y carrocería. Por eso solo son comparables las tarifas que el cliente
paga por viaje según el tipo de vehículo («POR TIPO DE VEHICULO TURBO»…). Las
tarifas por kilo, por caja o por metro cúbico no tienen contra qué compararse
sin saber cuánto lleva cada viaje: se informan como no comparables.

La categoría del cliente (TURBO, SENCILLO, TRACTOMULA…) se traduce a la
configuración SICETAC (2L3, 2, 3S2…) con el mapeo de vehículos de TarifaX, el
mismo que usa el motor de tarifas: una sola equivalencia para toda la
plataforma.

Variación = (tarifa del cliente − costo SICETAC) / costo SICETAC. Positiva: el
cliente paga por encima del costo de referencia. Negativa: por debajo.
"""
import io
import re
import unicodedata
from typing import Dict, List, Optional

import pandas as pd

# Columnas del consolidado comercial (los encabezados se comparan normalizados).
COLS_NEGOCIACION = {
    "CLIENTE": "cliente", "TIPO DE SERVICIO": "servicio", "ORIGEN": "origen",
    "DEPARTAMENTO ORIGEN": "dpto_origen", "DESTINO": "destino",
    "DEPARTAMENTO DESTINO": "dpto_destino", "TIPO DE TARIFA": "tipo_tarifa",
    "TARIFA": "tarifa", "BP": "bp",
}
OBLIGATORIAS_NEGOCIACION = ["CLIENTE", "ORIGEN", "DESTINO", "TIPO DE TARIFA", "TARIFA"]

# Columnas del archivo SICETAC, con los nombres alternos con que suele venir.
COLS_SICETAC = {
    "ORIGEN": ["ORIGEN", "MUNICIPIO ORIGEN", "CIUDAD ORIGEN"],
    "DESTINO": ["DESTINO", "MUNICIPIO DESTINO", "CIUDAD DESTINO"],
    "TIPO_VEHICULO": ["TIPO_VEHICULO", "TIPO VEHICULO", "CONFIGURACION", "CONFIGURACION VEHICULAR"],
    "TIPO_CARROCERIA": ["TIPO_CARROCERIA", "TIPO CARROCERIA", "CARROCERIA", "UNIDAD TRANSPORTE", "UNIDAD_TRANSPORTE"],
    "COSTO_TOTAL_VIAJE": ["COSTO_TOTAL_VIAJE", "COSTO TOTAL VIAJE", "COSTO TOTAL", "VALOR TOTAL VIAJE", "COSTO VIAJE"],
    "DISTANCIA": ["DISTANCIA", "DISTANCIA KM", "KILOMETROS", "KM"],
    # Fecha de publicación: si una ruta viene con varias, vale la más reciente.
    "FECHA_SICETAC": ["FECHA_SICETAC", "FECHA SICETAC", "PERIODO", "FECHA"],
    # Si el archivo no trae el departamento aparte, lo trae pegado a la ciudad
    # («RIOHACHA LA GUAJIRA») y hay que quitarlo (ver `_sin_departamento`).
    "DPTO_ORIGEN": ["DPTO_ORIGEN", "DPTO ORIGEN", "DEPARTAMENTO ORIGEN"],
    "DPTO_DESTINO": ["DPTO_DESTINO", "DPTO DESTINO", "DEPARTAMENTO DESTINO"],
    # Componentes del costo en el SICETAC oficial:
    #   COSTO_TOTAL_VIAJE = COSTO_MOVILIZACION + HORAS_LOGISTICAS × COSTO_HORA
    # Con ellos las horas logísticas del análisis REEMPLAZAN a las del archivo
    # en vez de sumarse encima (el total ya las trae).
    "COSTO_MOVILIZACION": ["COSTO_MOVILIZACION", "COSTO MOVILIZACION", "VALOR MOVILIZACION"],
    "HORAS_LOGISTICAS": ["HORAS_LOGISTICAS", "HORAS LOGISTICAS"],
    # General, granel sólido, refrigerada…: cambia el costo de la misma ruta.
    "TIPO_CARGA": ["TIPO_CARGA", "TIPO CARGA"],
    # Solo cuenta la vía estándar con el vehículo cargado (ver `_parcial`).
    "CONDICION_CARGA": ["CONDICION_CARGA", "CONDICION CARGA"],
    "VIA_ESTANDAR": ["VIA_ESTANDAR", "VIA ESTANDAR"],
    # Valor de UNA hora logística (espera en cargue y descargue) para esa
    # ruta y vehículo. Ojo: en el formato con «VALOR HORA», «COSTO HORA» es el
    # de todas las horas; por eso VALOR HORA va antes en la lista. Opcional: si el archivo no lo trae, se usa el valor por
    # configuración que se define en el tablero.
    "VALOR_HORA_LOGISTICA": ["VALOR_HORA_LOGISTICA", "VALOR HORA LOGISTICA", "HORA_LOGISTICA", "HORA LOGISTICA",
                             "VALOR_HORA", "VALOR HORA", "COSTO_HORA_LOGISTICA", "COSTO HORA LOGISTICA", "COSTO_HORA"],
}
OBLIGATORIAS_SICETAC = ["ORIGEN", "DESTINO", "TIPO_VEHICULO", "COSTO_TOTAL_VIAJE"]

# Sinónimos de categoría de vehículo que el mapeo puede tener con otro nombre.
SINONIMOS_CATEGORIA = {"TRACTOMULA": "TRACTOCAMION", "TRACTOCAMION": "TRACTOMULA", "CAMION SENCILLO": "SENCILLO"}

PREFIJO_VEHICULO = re.compile(r"^POR TIPO DE VEHICULO\s+", re.I)
# Nombres de vehículo que se reconocen dentro del tipo de tarifa aunque no
# venga con el prefijo «POR TIPO DE VEHICULO» («TARIFA TURBO», «MULA X VIAJE»).
# El orden importa: los compuestos antes que sus partes (TRACTOMULA antes que MULA).
VEHICULOS_CONOCIDOS = ["TRACTOCAMION", "TRACTOMULA", "MINIMULA", "DOBLETROQUE", "DOBLE TROQUE", "PATINETA",
                       "SENCILLO", "TURBO", "CAMIONETA", "CAMION", "MULA", "VANS", "NHR", "FURGON"]
# Lo que se paga por peso, caja o volumen no es por viaje, aunque nombre un vehículo.
NO_POR_VIAJE = re.compile(r"\b(KILO|KILOS|KG|CAJA|CAJAS|METRO|M3|UNIDAD|PALLET|ESTIBA)\b")

MESES = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto",
         "Septiembre", "Octubre", "Noviembre", "Diciembre"]


class ArchivoInvalido(ValueError):
    pass


def normalizar(v) -> str:
    """Texto comparable: mayúsculas, sin tildes, sin espacios dobles."""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    s = re.sub(r"\s+", " ", str(v).strip().upper())
    return unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")


def categoria_de(tipo_tarifa) -> Optional[str]:
    """El tipo de vehículo que trae el tipo de tarifa, si es una tarifa por viaje.

    «POR TIPO DE VEHICULO TURBO» → «TURBO»; «TARIFA TURBO» → «TURBO».
    «TARIFA POR KILO» o «POR METRO CUBICO 0,5 A 5» → None: no son por viaje.
    """
    t = normalizar(tipo_tarifa)
    if not t:
        return None
    if t.startswith("VIAJE "):          # Consolidado SG: «VIAJE TURBO SENCILLO», «VIAJE TURBO <45»
        return t[6:].strip() or None
    if PREFIJO_VEHICULO.match(t):
        return PREFIJO_VEHICULO.sub("", t).strip() or None
    if NO_POR_VIAJE.search(t):
        return None
    for v in VEHICULOS_CONOCIDOS:
        if re.search(r"\b" + v + r"\b", t):
            return v.replace(" ", "")
    return None


def _leer_tabla(contenido: bytes, nombre: str) -> pd.DataFrame:
    try:
        if nombre.lower().endswith(".csv"):
            return pd.read_csv(io.BytesIO(contenido), sep=None, engine="python", encoding_errors="replace")
        libro = pd.ExcelFile(io.BytesIO(contenido))
        # Un libro con varias hojas trae las tarifas depuradas en la que empieza
        # por «TARIFAS_» (p. ej. TARIFAS_FINALES); si no hay, la primera.
        hoja = next((h for h in libro.sheet_names if normalizar(h).startswith("TARIFAS_")), libro.sheet_names[0])
        return pd.read_excel(libro, sheet_name=hoja)
    except Exception as e:  # noqa: BLE001
        raise ArchivoInvalido(f"No se pudo leer el archivo «{nombre}»: {e}") from e


# ── Consolidado SG (formato de las comunicaciones comerciales) ──
# Columnas: N° SG, Cliente, CEDI, Servicio Principal, Clasificación (vehículo),
# Descripción, «Origen – Destino» (un solo texto), Medida, Valor Unitario,
# Vigencia Inicio, Vigencia Fin…
_SEP_OD = re.compile(r"\s+[–—-]\s+")
_FIN_CIUDAD = r"(?=\s+(?:A|TARIFA|EN|CON|HACIA|HASTA|PARA|Y)\b|[:,.;(]|$)"
_DESDE_EN = re.compile(r"\b(?:DESDE|ORIGEN)\b[^:|]*?\bEN\s+([A-Z][A-Z ]*?)" + _FIN_CIUDAD)   # «DESDE EL COL UNILEVER EN PALMIRA A:»
_DESDE = re.compile(r"\b(?:DESDE|ORIGEN)\s+([A-Z][A-Z ]*?)" + _FIN_CIUDAD)
# Lo que sigue a «desde» y no es una ciudad: «las plataformas de Icoltrans», «el CEDI».
_NO_CIUDAD = re.compile(r"^(?:(?:EL|LA|LOS|LAS)\s+)?(?:PLATAFORMA|PLATAFORMAS|CEDI|BODEGA|BODEGAS|PLANTA|FABRICA|COL)\b")


def _origen_en_descripcion(descripcion) -> Optional[str]:
    texto = normalizar(descripcion)
    for patron in (_DESDE_EN, _DESDE):
        m = patron.search(texto)
        if m and not _NO_CIUDAD.match(m.group(1).strip()):
            return m.group(1).strip()
    return None


def _col(df: pd.DataFrame, *opciones: str) -> Optional[str]:
    """La columna cuyo nombre (sin tildes ni signos) empieza por alguna opción."""
    for c in df.columns:
        n = " ".join(re.sub(r"[^A-Z0-9 ]", " ", normalizar(c)).split())
        if any(n == o or n.startswith(o + " ") for o in opciones):
            return c
    return None


def es_consolidado_sg(df: pd.DataFrame) -> bool:
    """El Consolidado SG trae «Origen – Destino» en una columna; el informe de
    rutas (mismo formato), Origen y Destino por separado."""
    ruta = _col(df, "ORIGEN DESTINO") or (_col(df, "ORIGEN") and _col(df, "DESTINO"))
    return bool(ruta) and all(_col(df, x) for x in ("VALOR UNITARIO", "MEDIDA", "CLIENTE"))


def limpiar_ciudad(texto) -> Optional[str]:
    """Una columna de ciudad del informe de rutas, limpia:
    «BARRANQUILLA-ATLANTICO» → BARRANQUILLA · «Bogotá (Aviatur)» → BOGOTA ·
    «EL COL UNILEVER EN PALMIRA» → PALMIRA · «LA PLATAFORMA DE YUMBO» → YUMBO.
    Lo que no nombra una ciudad («LAS PLATAFORMAS DE ICOLTRANS») → None."""
    if texto is None or (isinstance(texto, float) and pd.isna(texto)):
        return None
    t = normalizar(re.sub(r"\(.*?\)", "", str(texto)))
    if not t:
        return None
    m = re.search(r"\bEN\s+([A-Z][A-Z ]*)$", t)            # «… EN PALMIRA», «… PLANTA DE JGB EN CALI»
    if m:
        t = m.group(1).strip()
    if "-" in t:                                            # «CALI-VALLE DEL CAUCA», «CALI-VALLE»
        a, b = [x.strip() for x in t.split("-", 1)]
        # «CALI-VALLE», «RIOHACHA-GUAJIRA»: el departamento, aunque venga abreviado.
        if not b or b in DEPARTAMENTOS or any(d.startswith(b + " ") or d.endswith(" " + b) for d in DEPARTAMENTOS):
            t = a
    if _NO_CIUDAD.match(t):                                 # «LA PLATAFORMA DE YUMBO»
        m = re.search(r"\bDE\s+([A-Z][A-Z ]*)$", t)
        t = m.group(1).strip() if m else ""
        if t in ("ICOLTRANS", ""):
            return None
    return t or None


def partir_origen_destino(od, cedi=None, descripcion=None):
    """«YUMBO – BOGOTÁ (CUNDINAMARCA)» → (YUMBO, BOGOTA, None, CUNDINAMARCA).

    Quita departamentos y zonas («4. Cundi- Boyacence», «ZONA 7»). Si solo
    queda la ciudad destino («Armenia – Quindío»), el origen sale de la
    descripción («… DESDE FUNZA A: …», «con origen Yumbo») o del CEDI."""
    if od is None or (isinstance(od, float) and pd.isna(od)) or not str(od).strip():
        return None, None, None, None
    deps = set(DEPARTAMENTOS) | {"VALLE", "NORTE DE S", "D.C.", "BOGOTA D.C"}
    ciudades, dpto_o, dpto_d, ciudad_dep = [], None, None, None
    for crudo in _SEP_OD.split(str(od)):
        entre = re.findall(r"\((.*?)\)", crudo)
        if entre:
            dpto_d = normalizar(entre[-1])
        t = normalizar(re.sub(r"\(.*?\)", "", crudo))
        if "-" in t:                                   # «BARRANQUILLA-ATLANTICO»
            a, b = [x.strip() for x in t.split("-", 1)]
            if not b or b in deps:
                t = a
                if b and not ciudades:
                    dpto_o = b
        if not t or re.match(r"^(\d|ZONA\b)", t):
            continue
        if t in deps:
            dpto_d = t
            ciudad_dep = ciudad_dep or t                # «CALDAS – ANTIOQUIA»: Caldas es municipio
            continue
        ciudades.append(t)
    if not ciudades and ciudad_dep:
        ciudades = [ciudad_dep]
    if len(ciudades) >= 2:
        return ciudades[0], ciudades[-1], dpto_o, dpto_d
    if len(ciudades) == 1:
        origen = (_origen_en_descripcion(descripcion) if descripcion is not None else None) or normalizar(cedi) or None
        return origen, ciudades[0], None, dpto_d
    return None, None, None, None


def _mes(v) -> Optional[str]:
    f = pd.to_datetime(v, errors="coerce", dayfirst=True)
    return None if pd.isna(f) else f.strftime("%Y-%m")


def leer_consolidado_sg(df: pd.DataFrame) -> pd.DataFrame:
    col = {k: _col(df, *o) for k, o in {
        "sg": ("N SG", "NO SG", "SG"), "cliente": ("CLIENTE",), "cedi": ("CEDI",), "servicio": ("SERVICIO PRINCIPAL", "SERVICIO"),
        "clasificacion": ("CLASIFICACION",), "descripcion": ("DESCRIPCION",), "od": ("ORIGEN DESTINO", "ORIGEN"),
        "medida": ("MEDIDA",), "valor": ("VALOR UNITARIO",), "desde": ("VIGENCIA INICIO",), "hasta": ("VIGENCIA FIN",),
    }.items()}
    get = lambda k: df[col[k]] if col[k] else pd.Series([None] * len(df), index=df.index)  # noqa: E731
    if col["od"] and normalizar(col["od"]).replace(" ", "").startswith("ORIGEN") and "DESTINO" in normalizar(col["od"]):
        partes = [partir_origen_destino(od, cedi, desc) for od, cedi, desc in zip(get("od"), get("cedi"), get("descripcion"))]
    else:
        # Informe de rutas: Origen y Destino por separado. Sin origen escrito,
        # sale de la descripción («DESDE FUNZA A:») o del CEDI.
        def dpto(v):
            n = normalizar(v)
            return n.split("-", 1)[1].strip() or None if "-" in n else None
        partes = []
        for o, d, cedi, desc in zip(df[_col(df, "ORIGEN")], df[_col(df, "DESTINO")], get("cedi"), get("descripcion")):
            origen = limpiar_ciudad(o) or (_origen_en_descripcion(desc) if desc is not None else None) or normalizar(cedi) or None
            partes.append((origen, limpiar_ciudad(d), dpto(o), dpto(d)))
    medida, clase = get("medida").map(normalizar), get("clasificacion").map(normalizar)
    # Tipo de tarifa tal como viene: la medida, y el vehículo si es por viaje
    # («VIAJE TURBO», «KILO», «CAJA»). Así el filtro separa paqueteo de masivo
    # y la categoría se extrae igual que en el formato anterior.
    tipo = [f"{m} {c}".strip() if m == "VIAJE" else m for m, c in zip(medida, clase)]
    out = pd.DataFrame({
        "bp": get("sg").map(lambda v: normalizar(v)[:40] or None), "cliente": get("cliente").map(lambda v: normalizar(v) or None),
        "servicio": get("servicio").map(lambda v: normalizar(v) or None),
        "origen": [p[0] for p in partes], "destino": [p[1] for p in partes],
        "dpto_origen": [p[2] for p in partes], "dpto_destino": [p[3] for p in partes],
        "tipo_tarifa": [t or None for t in tipo], "tarifa": pd.to_numeric(get("valor"), errors="coerce"),
        "vigencia_desde": get("desde").map(_mes), "vigencia_hasta": get("hasta").map(_mes),
    }, index=df.index)
    out = out[out["cliente"].notna() & out["origen"].notna() & out["destino"].notna()]
    # Una «ruta» de más de 60 letras es una nota escrita en esa columna, no una ciudad.
    out = out[(out["origen"].str.len() <= 60) & (out["destino"].str.len() <= 60)]
    if out.empty:
        raise ArchivoInvalido("El consolidado no trae ninguna fila con cliente, origen y destino.")
    out["categoria_vehiculo"] = out["tipo_tarifa"].map(categoria_de)
    return out.reset_index(drop=True)


# Largo máximo de cada texto en la base (tarifax_negociaciones).
LARGOS = {"cliente": 200, "servicio": 120, "origen": 120, "destino": 120, "dpto_origen": 120, "dpto_destino": 120,
          "tipo_tarifa": 120, "bp": 40}


def _recortar(df: pd.DataFrame) -> pd.DataFrame:
    for c, n in LARGOS.items():
        if c in df:
            df[c] = df[c].map(lambda v: v[:n] if isinstance(v, str) else v)
    return df


def leer_negociaciones(contenido: bytes, nombre: str) -> pd.DataFrame:
    return _recortar(_leer_negociaciones(contenido, nombre))


def _leer_negociaciones(contenido: bytes, nombre: str) -> pd.DataFrame:
    df = _leer_tabla(contenido, nombre)
    if es_consolidado_sg(df):
        return leer_consolidado_sg(df)
    encabezados = {normalizar(c): c for c in df.columns}
    faltan = [c for c in OBLIGATORIAS_NEGOCIACION if c not in encabezados]
    if faltan:
        raise ArchivoInvalido(
            "Al archivo de negociaciones le faltan las columnas: " + ", ".join(faltan)
            + ". Se esperan: " + ", ".join(COLS_NEGOCIACION) + ".")
    out = pd.DataFrame({dest: df[encabezados[src]] if src in encabezados else None
                        for src, dest in COLS_NEGOCIACION.items()})
    for c in ("cliente", "servicio", "origen", "dpto_origen", "destino", "dpto_destino", "tipo_tarifa", "bp"):
        out[c] = out[c].map(lambda v: normalizar(v) or None)
    out["tarifa"] = pd.to_numeric(out["tarifa"], errors="coerce")
    out = out[out["cliente"].notna() & out["origen"].notna() & out["destino"].notna()]
    if out.empty:
        raise ArchivoInvalido("El archivo de negociaciones no trae ninguna fila con cliente, origen y destino.")
    out["categoria_vehiculo"] = out["tipo_tarifa"].map(categoria_de)
    # AÑO y MES (opcionales, de la plantilla por periodo): desde cuándo rige cada fila.
    col_anio = next((encabezados[c] for c in ("ANO", "AÑO", "ANIO") if c in encabezados), None)
    col_mes = encabezados.get("MES")
    if col_anio is not None and col_mes is not None:
        anio = pd.to_numeric(df.loc[out.index, col_anio], errors="coerce")
        mes = pd.to_numeric(df.loc[out.index, col_mes], errors="coerce")
        hay = anio.notna() | mes.notna()
        malas = hay & ~(anio.between(2000, 2100) & mes.between(1, 12))
        if malas.any():
            filas = ", ".join(str(int(i) + 2) for i in out.index[malas][:10])
            raise ArchivoInvalido(f"AÑO o MES inválido en las filas {filas} del archivo (MES de 1 a 12).")
        out["periodo"] = [f"{int(a)}-{int(m):02d}" if h else None for a, m, h in zip(anio, mes, hay)]
    return out.reset_index(drop=True)


SIN_PERIODO = "0000-00"


def vigentes(negociaciones: List[dict], periodo: str, hoy: Optional[str] = None) -> List[dict]:
    """Las tarifas que rigen en un mes.

    - Un cliente con cargas por periodo usa su carga más reciente de ese mes o
      antes (reemplaza todas sus tarifas).
    - Si no tiene, rige el consolidado general: las filas cuya vigencia
      (desde/hasta, vacías = sin límite) incluye el mes; si una misma tarifa
      tiene dos vigencias que se cruzan, la que empezó después.
    «VIGENTE» = el mes en curso."""
    if periodo == "VIGENTE":
        from datetime import date
        periodo = hoy or date.today().strftime("%Y-%m")
    elegido: Dict[str, str] = {}
    for n in negociaciones:
        p = n.get("periodo")
        if p and p <= periodo and p > elegido.get(n["cliente"], ""):
            elegido[n["cliente"]] = p
    out = [n for n in negociaciones if n.get("periodo") and elegido.get(n["cliente"]) == n["periodo"]]
    mejores: Dict[str, dict] = {}
    for n in negociaciones:
        if n.get("periodo") or n["cliente"] in elegido:
            continue
        d, h = n.get("vigencia_desde"), n.get("vigencia_hasta")
        if (d and d > periodo) or (h and h < periodo):
            continue
        k = llave_tarifa(n["cliente"], n.get("servicio"), n.get("origen"), n.get("destino"), n.get("tipo_tarifa"))
        if k not in mejores or (d or "") >= (mejores[k].get("vigencia_desde") or ""):
            mejores[k] = n
    return out + list(mejores.values())


def llave_tarifa(cliente, servicio, origen, destino, tipo_tarifa) -> str:
    """La misma tarifa en distintos periodos (aunque cambie su valor)."""
    return "|".join(str(x or "") for x in (cliente, servicio, origen, destino, tipo_tarifa))


def _columnas_sicetac(encabezados) -> Dict[str, int]:
    """Posición de cada columna que usa el cruce, reconociendo sus nombres alternos."""
    pos: Dict[str, int] = {}
    vistos = {}
    for i, c in enumerate(encabezados):
        n = normalizar(c)
        vistos.setdefault(n, i)
        vistos.setdefault(n.replace(" ", "_"), i)
    for destino, alternos in COLS_SICETAC.items():
        for alt in alternos:
            if alt in vistos:
                pos[destino] = vistos[alt]
                break
    return pos


def _falta_sicetac(pos: Dict[str, int], encontradas=None) -> None:
    faltan = [c for c in OBLIGATORIAS_SICETAC if c not in pos]
    if faltan:
        vistas = [str(c) for c in (encontradas or []) if c not in (None, "")]
        raise ArchivoInvalido(
            "Al archivo SICETAC le faltan las columnas: " + ", ".join(faltan)
            + ". Debe traer ORIGEN, DESTINO, TIPO_VEHICULO, TIPO_CARROCERIA y COSTO_TOTAL_VIAJE"
              " (y, si las tiene, COSTO_MOVILIZACION, COSTO_HORA, HORAS_LOGISTICAS, TIPO_CARGA)."
            + (f" Columnas encontradas en el archivo: {', '.join(vistas[:30])}." if vistas else
               " No se encontró una fila de encabezados en las primeras 50 filas."))


def _norm_rapida(s: pd.Series) -> pd.Series:
    """`normalizar` sobre cada valor DISTINTO y no sobre cada fila: en un
    tarifario de millones de filas las ciudades y vehículos se repiten mucho."""
    unicos = pd.unique(s)
    return s.map({u: normalizar(u) for u in unicos})


LLAVES = ["o", "d", "tv", "carroceria", "tipo_carga"]

DEPARTAMENTOS = sorted([
    "AMAZONAS", "ANTIOQUIA", "ARAUCA", "ATLANTICO", "BOLIVAR", "BOYACA", "CALDAS", "CAQUETA", "CASANARE", "CAUCA",
    "CESAR", "CHOCO", "CORDOBA", "CUNDINAMARCA", "GUAINIA", "GUAVIARE", "HUILA", "LA GUAJIRA", "MAGDALENA", "META",
    "NARINO", "NORTE DE SANTANDER", "PUTUMAYO", "QUINDIO", "RISARALDA", "SANTANDER", "SUCRE", "TOLIMA",
    "VALLE DEL CAUCA", "VAUPES", "VICHADA", "BOGOTA D. C.", "BOGOTA D.C.", "BOGOTA, D.C.", "BOGOTA D C", "BOGOTA, DISTRITO CAPITAL", "BOGOTA DISTRITO CAPITAL",
    "SAN ANDRES Y PROVIDENCIA", "ARCHIPIELAGO DE SAN ANDRES, PROVIDENCIA Y SANTA CATALINA", "SAN ANDRES",
], key=len, reverse=True)   # los largos primero: VALLE DEL CAUCA antes que CAUCA


def _sin_departamento(ciudad: str) -> str:
    """«RIOHACHA LA GUAJIRA» → «RIOHACHA»; «BOGOTA BOGOTA D. C.» → «BOGOTA».
    Solo se quita un departamento, y solo si queda algo antes."""
    for dep in DEPARTAMENTOS:
        if ciudad.endswith(" " + dep):
            return ciudad[: -len(dep) - 1].strip()
    return ciudad


def _parcial(base: pd.DataFrame) -> pd.DataFrame:
    """Sumas y conteos por ruta, vehículo, carrocería y tipo de carga de un
    bloque de filas. Varios parciales se combinan en `_combinar`.

    Del SICETAC oficial se toma solo la vía estándar con el vehículo cargado:
    las vías alternas y los retornos vacíos son otros costos, y promediarlos
    con el principal falsearía la comparación."""
    if "VIA_ESTANDAR" in base:
        via = _norm_rapida(base["VIA_ESTANDAR"])
        base = base[via.isin(["SI", "S", "TRUE", "1", "VERDADERO", ""])]
    if "CONDICION_CARGA" in base:
        cond = _norm_rapida(base["CONDICION_CARGA"])
        base = base[cond.isin(["CARGADO", ""])]
    n = len(base)
    vacio = pd.Series([float("nan")] * n, index=base.index)
    if "FECHA_SICETAC" in base:
        unicos = pd.unique(base["FECHA_SICETAC"])
        legibles = pd.to_datetime(pd.Series(unicos), errors="coerce").dt.strftime("%Y-%m-%d").fillna("")
        fecha = base["FECHA_SICETAC"].map(dict(zip(unicos, legibles)))
    else:
        fecha = ""
    o, d = _norm_rapida(base["ORIGEN"]), _norm_rapida(base["DESTINO"])
    if "DPTO_ORIGEN" not in base:
        o = o.map({u: _sin_departamento(u) for u in pd.unique(o)})
    if "DPTO_DESTINO" not in base:
        d = d.map({u: _sin_departamento(u) for u in pd.unique(d)})
    g = pd.DataFrame({
        "o": o,
        "d": d,
        "fecha": fecha,
        "tv": _norm_rapida(base["TIPO_VEHICULO"]),
        "carroceria": _norm_rapida(base["TIPO_CARROCERIA"]) if "TIPO_CARROCERIA" in base else "SIN CARROCERIA",
        "tipo_carga": _norm_rapida(base["TIPO_CARGA"]) if "TIPO_CARGA" in base else "",
        "costo": pd.to_numeric(base["COSTO_TOTAL_VIAJE"], errors="coerce"),
        "mov": pd.to_numeric(base["COSTO_MOVILIZACION"], errors="coerce") if "COSTO_MOVILIZACION" in base else vacio,
        "horas": pd.to_numeric(base["HORAS_LOGISTICAS"], errors="coerce") if "HORAS_LOGISTICAS" in base else vacio,
        "distancia": pd.to_numeric(base["DISTANCIA"], errors="coerce") if "DISTANCIA" in base else vacio,
        "valor_hora": pd.to_numeric(base["VALOR_HORA_LOGISTICA"], errors="coerce")
                      if "VALOR_HORA_LOGISTICA" in base else vacio,
    })
    g["carroceria"] = g["carroceria"].replace("", "SIN CARROCERIA")
    g = g[(g["o"] != "") & (g["d"] != "") & (g["tv"] != "") & g["costo"].gt(0)]
    return g.groupby(LLAVES + ["fecha"], as_index=False).agg(
        costo_suma=("costo", "sum"), costo_n=("costo", "count"),
        mov_suma=("mov", "sum"), mov_n=("mov", "count"),
        horas_suma=("horas", "sum"), horas_n=("horas", "count"),
        hora_suma=("valor_hora", "sum"), hora_n=("valor_hora", "count"),
        distancia=("distancia", "first"))


def _combinar(parciales: List[pd.DataFrame]) -> pd.DataFrame:
    if not parciales:
        raise ArchivoInvalido("El archivo SICETAC no trae ninguna fila con origen, destino, vehículo y costo.")
    t = pd.concat(parciales, ignore_index=True).groupby(LLAVES + ["fecha"], as_index=False).agg(
        costo_suma=("costo_suma", "sum"), costo_n=("costo_n", "sum"), mov_suma=("mov_suma", "sum"),
        mov_n=("mov_n", "sum"), horas_suma=("horas_suma", "sum"), horas_n=("horas_n", "sum"),
        hora_suma=("hora_suma", "sum"), hora_n=("hora_n", "sum"), distancia=("distancia", "first"))
    t = t[t["costo_n"] > 0]
    if t.empty:
        raise ArchivoInvalido("El archivo SICETAC no trae ninguna fila de vía estándar, vehículo cargado y costo.")
    # Si la misma ruta viene publicada en varias fechas, vale la más reciente
    # (no el promedio entre fechas). Sin fecha («») cuenta como la más antigua.
    t = t[t["fecha"] == t.groupby(LLAVES)["fecha"].transform("max")]
    t["costo"] = t["costo_suma"] / t["costo_n"]
    t["costo_mov"] = (t["mov_suma"] / t["mov_n"]).where(t["mov_n"] > 0)
    t["horas_sicetac"] = (t["horas_suma"] / t["horas_n"]).where(t["horas_n"] > 0)
    t["valor_hora"] = (t["hora_suma"] / t["hora_n"]).where(t["hora_n"] > 0)
    return t[LLAVES + ["fecha", "costo", "costo_mov", "horas_sicetac", "valor_hora", "distancia"]].reset_index(drop=True)


def agrupar_sicetac(df: pd.DataFrame) -> pd.DataFrame:
    """Costo promedio, componentes y distancia por ruta, vehículo, carrocería y tipo de carga."""
    pos = _columnas_sicetac(df.columns)
    _falta_sicetac(pos, list(df.columns))
    return _combinar([_parcial(pd.DataFrame({k: df.iloc[:, i] for k, i in pos.items()}))])


BLOQUE = 100_000


def leer_sicetac_archivo(ruta: str, nombre: str, avance=None) -> pd.DataFrame:
    """Lee un tarifario SICETAC de cualquier tamaño sin cargarlo entero en
    memoria: Excel fila a fila (modo de solo lectura), CSV por bloques. Cada
    bloque se agrupa y se suma; la memoria depende de cuántas rutas distintas
    hay, no de cuántas filas trae el archivo.

    (Se probó un lector de Excel 10 veces más rápido, pero carga el archivo
    entero: 1,6 GB por un SICETAC mensual. En un servidor que ya pierde
    procesos por memoria, se prefiere lento y estable.)

    `avance(filas_leidas, filas_totales)` se llama después de cada bloque."""
    parciales: List[pd.DataFrame] = []
    leidas = 0
    nombre = (nombre or "").lower()
    try:
        if nombre.endswith(".csv"):
            with open(ruta, "r", encoding="utf-8", errors="replace") as fh:
                primera = fh.readline()
            sep = max([";", ",", "\t", "|"], key=primera.count)
            pos = None
            for bloque in pd.read_csv(ruta, sep=sep, dtype=str, chunksize=BLOQUE, encoding_errors="replace"):
                if pos is None:
                    pos = _columnas_sicetac(bloque.columns)
                    _falta_sicetac(pos, list(bloque.columns))
                base = pd.DataFrame({k: bloque.iloc[:, i] for k, i in pos.items()})
                parciales.append(_parcial(base))
                leidas += len(bloque)
                if avance:
                    avance(leidas, None)
        elif nombre.endswith(".xlsx"):
            import openpyxl
            libro = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
            # Una hoja de Excel no pasa de 1.048.576 filas: los tarifarios grandes
            # pueden venir en varias hojas. Se leen todas las que tengan las
            # columnas del tarifario (y se ignoran las de notas o resúmenes).
            total = sum((h.max_row or 0) for h in libro.worksheets) or None
            hojas_leidas = 0
            vistas: List = []
            for hoja in libro.worksheets:
                filas = hoja.iter_rows(values_only=True)
                pos = None
                # El encabezado puede no estar en la primera fila (títulos arriba).
                for _ in range(50):
                    fila = next(filas, None)
                    if fila is None:
                        break
                    if any(c not in (None, "") for c in fila) and not vistas:
                        vistas = list(fila)
                    candidato = _columnas_sicetac([c if c is not None else "" for c in fila])
                    if all(c in candidato for c in OBLIGATORIAS_SICETAC):
                        pos = candidato
                        break
                if pos is None:
                    continue
                hojas_leidas += 1
                claves = list(pos.items())
                lote: List[tuple] = []
                for fila in filas:
                    lote.append(tuple(fila[i] if i < len(fila) else None for _, i in claves))
                    if len(lote) >= BLOQUE:
                        parciales.append(_parcial(pd.DataFrame(lote, columns=[k for k, _ in claves])))
                        leidas += len(lote)
                        lote = []
                        if avance:
                            avance(leidas, total)
                if lote:
                    parciales.append(_parcial(pd.DataFrame(lote, columns=[k for k, _ in claves])))
                    leidas += len(lote)
                    if avance:
                        avance(leidas, total)
            libro.close()
            if not hojas_leidas:
                _falta_sicetac({}, vistas)
        else:
            # .xls (formato viejo): no se puede leer por partes; se lee entero.
            with open(ruta, "rb") as fh:
                df = _leer_tabla(fh.read(), nombre)
            pos = _columnas_sicetac(df.columns)
            _falta_sicetac(pos, list(df.columns))
            parciales.append(_parcial(pd.DataFrame({k: df.iloc[:, i] for k, i in pos.items()})))
            leidas = len(df)
    except ArchivoInvalido:
        raise
    except Exception as e:  # noqa: BLE001
        raise ArchivoInvalido(f"No se pudo leer el archivo SICETAC: {e}") from e
    if avance:
        avance(leidas, leidas)
    return _combinar(parciales)


def leer_sicetac(contenido: bytes, nombre: str) -> pd.DataFrame:
    """Lo mismo, desde un archivo ya en memoria (archivos pequeños y pruebas)."""
    import os
    import tempfile
    sufijo = os.path.splitext(nombre or "")[1] or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=sufijo, delete=False) as tmp:
        tmp.write(contenido)
        ruta = tmp.name
    try:
        return leer_sicetac_archivo(ruta, nombre or f"x{sufijo}")
    finally:
        os.unlink(ruta)


def codigo_vehiculo(v) -> str:
    """La configuración sin su descripción: «2L2 Liviano entre 8 y 9 Tonel.» →
    «2L2». Cada formato de SICETAC la escribe distinto; el código es el mismo."""
    n = normalizar(v)
    return n.split(" ")[0] if n else ""


def mapeo_categorias(mapeo: Dict[str, str]) -> Dict[str, str]:
    """{CATEGORIA_NORMALIZADA: código de configuración SICETAC}, con sinónimos."""
    out = {normalizar(k): codigo_vehiculo(v) for k, v in mapeo.items() if normalizar(k) and codigo_vehiculo(v)}
    for a, b in SINONIMOS_CATEGORIA.items():
        if a not in out and b in out:
            out[a] = out[b]
    return out


def cruzar(negociaciones: List[dict], sicetac: pd.DataFrame, alias: Dict[str, str],
           mapeo: Dict[str, str], periodo: str) -> List[dict]:
    """Una fila por tarifa comparable, carrocería y tipo de carga con referencia SICETAC; si la
    ruta no tiene referencia, una sola fila marcada SIN_RUTA (o SIN_VEHICULO si
    la categoría no tiene equivalente en el mapeo)."""
    alias_n = {normalizar(k): normalizar(v) for k, v in alias.items()}
    mapa = mapeo_categorias(mapeo)
    indice: Dict[tuple, List[tuple]] = {}
    # Los archivos guardados con versiones anteriores no traen todas las columnas.
    for col, defecto in (("valor_hora", float("nan")), ("tipo_carga", ""), ("costo_mov", float("nan")),
                         ("horas_sicetac", float("nan")), ("fecha", "")):
        if col not in sicetac:
            sicetac = sicetac.assign(**{col: defecto})
    # La equivalencia de ciudades se aplica a los dos lados: un formato de
    # SICETAC dice «BOGOTA, D.C.» y otro «BOGOTA», igual que las negociaciones.
    ciudad = lambda c: alias_n.get(c, c)  # noqa: E731
    # Solo las filas de SICETAC con un origen y un destino que alguna tarifa usa:
    # un mes trae más de un millón de filas y las negociaciones, pocas rutas.
    usadas_o = {alias_n.get(normalizar(n["origen"]), normalizar(n["origen"])) for n in negociaciones}
    usadas_d = {alias_n.get(normalizar(n["destino"]), normalizar(n["destino"])) for n in negociaciones}
    canon_o = {c: ciudad(c) for c in pd.unique(sicetac["o"])}
    canon_d = {c: ciudad(c) for c in pd.unique(sicetac["d"])}
    sicetac = sicetac[sicetac["o"].map(canon_o).isin(usadas_o) & sicetac["d"].map(canon_d).isin(usadas_d)]
    codigos = {t: codigo_vehiculo(t) for t in pd.unique(sicetac["tv"])}
    columnas = ["o", "d", "tv", "carroceria", "tipo_carga", "costo", "distancia", "valor_hora", "costo_mov",
                "horas_sicetac", "fecha"]
    for o, d, tv, *resto in sicetac[columnas].itertuples(index=False):
        indice.setdefault((ciudad(o), ciudad(d), codigos[tv]), []).append(tuple(resto))
    filas = []
    for n in negociaciones:
        cat = n["categoria_vehiculo"]
        # Desde pandas lo vacío llega como NaN, que es «verdadero».
        if not isinstance(cat, str) or not cat:
            continue
        # Una tarifa vacía o en cero no es una negociación: compararla daría
        # «−100 % frente a SICETAC», una pérdida que no existe.
        if n["tarifa"] is None or pd.isna(n["tarifa"]) or n["tarifa"] <= 0:
            continue
        o = alias_n.get(normalizar(n["origen"]), normalizar(n["origen"]))
        d = alias_n.get(normalizar(n["destino"]), normalizar(n["destino"]))
        tv = mapa.get(normalizar(cat))
        comun = dict(periodo=periodo, negociacion_id=n["id"], cliente=n["cliente"], servicio=n["servicio"],
                     origen=n["origen"], destino=n["destino"], tipo_tarifa=n.get("tipo_tarifa"),
                     dpto_origen=n.get("dpto_origen"),
                     dpto_destino=n.get("dpto_destino"), categoria=cat, tipo_vehiculo=tv,
                     tarifa_cliente=n["tarifa"], valor_hora=None, tipo_carga=None, costo_mov=None, horas_sicetac=None,
                     fecha_sicetac=None)
        if not tv:
            filas.append({**comun, "estado": "SIN_VEHICULO", "carroceria": None, "distancia": None,
                          "costo_sicetac": None, "diferencia": None, "variacion_pct": None})
            continue
        refs = indice.get((o, d, tv))
        if not refs:
            filas.append({**comun, "estado": "SIN_RUTA", "carroceria": None, "distancia": None,
                          "costo_sicetac": None, "diferencia": None, "variacion_pct": None})
            continue
        for carr, carga, costo, dist, vh, mov, hs, fecha in refs:
            tarifa = n["tarifa"]
            dif = (tarifa - costo) if tarifa is not None else None
            filas.append({**comun, "estado": "OK", "carroceria": carr, "tipo_carga": carga or None,
                          "fecha_sicetac": fecha or None,
                          "valor_hora": None if pd.isna(vh) else round(float(vh), 0),
                          "costo_mov": None if pd.isna(mov) else round(float(mov), 0),
                          "horas_sicetac": None if pd.isna(hs) else round(float(hs), 2),
                          "distancia": None if pd.isna(dist) else float(dist),
                          "costo_sicetac": round(float(costo), 0),
                          "diferencia": None if dif is None else round(float(dif), 0),
                          "variacion_pct": None if dif is None or not costo else round(float(dif) / float(costo) * 100, 2)})
    return filas


def nombre_periodo(periodo: str) -> str:
    if periodo == "VIGENTE":
        return "Tarifario TarifaX vigente"
    anio, mes = periodo.split("-")
    return f"{MESES[int(mes) - 1]} {anio}"
